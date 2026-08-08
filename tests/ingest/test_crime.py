"""Tests for GUS Panorama dzielnic crime ingest and listing join."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from happy_housing.ingest.crime import (
    CRIME_FEATURE_COLUMNS,
    attach_crime_features,
    build_crime_by_district,
    load_crime_by_district,
    parse_panorama_crime_tables,
)


def _write_minimal_panorama(tmp_path: Path) -> tuple[Path, Path]:
    """Build tiny Bezpieczeństwo + Ludność workbooks matching Panorama layout."""
    from openpyxl import Workbook

    safe = tmp_path / "04_Bezpieczenstwo.xlsx"
    wb = Workbook()
    ws16 = wb.active
    ws16.title = "16"
    ws16.append(["title"])
    ws16.append(["en"])
    ws16.append([None])
    ws16.append(["Wyszczególnienie", "Ogółem", "kryminalne", "gosp", "drog", "inne"])
    ws16.append([None, None, None, None, None, None])
    ws16.append(["M.st. Warszawa a", 100, 60, 20, 10, 10])
    ws16.append(["Warsaw a", None, None, None, None, None])
    ws16.append(["Mokotów", 50, 30, 10, 5, 5])
    ws16.append(["Wola", 40, 25, 8, 4, 3])
    ws17 = wb.create_sheet("17")
    ws17.append(["title"])
    ws17.append(["en"])
    ws17.append([None])
    ws17.append(["Wyszczególnienie", "Ogółem", "krym", "gosp", "drog"])
    ws17.append([None, "w %", None, None, None])
    ws17.append(["M.st. Warszawa a", 50.0, 40.0, 55.0, 96.0])
    ws17.append(["Warsaw a", None, None, None, None])
    ws17.append(["Mokotów", 54.3, 46.0, 66.0, 94.0])
    ws17.append(["Wola", 46.5, 45.0, 43.0, 94.0])
    wb.save(safe)

    pop = tmp_path / "05_Ludnosc.xlsx"
    wb2 = Workbook()
    ws23 = wb2.active
    ws23.title = "23"
    ws23.append(["title"])
    ws23.append(["en"])
    ws23.append([None])
    ws23.append(["Dzielnice", "Ogółem", None, None, None, None])
    ws23.append([None, None, None, None, None, None])
    ws23.append([None, "30.06", "31.12", None, None, None])
    ws23.append(["M.st. Warszawa", 2000, 2000, None, None, None])
    ws23.append(["Warsaw", None, None, None, None, None])
    ws23.append(["Mokotów", 900, 1000, None, None, None])
    ws23.append(["Wola", 700, 800, None, None, None])
    wb2.save(pop)
    return safe, pop


def test_parse_panorama_crime_tables_builds_rates(tmp_path: Path) -> None:
    safe, pop = _write_minimal_panorama(tmp_path)
    frame = parse_panorama_crime_tables(safe, pop, year=2024)
    assert set(frame["district"]) == {"Mokotów", "Wola"}
    mok = frame.loc[frame["district"] == "Mokotów"].iloc[0]
    assert int(mok["crime_total"]) == 50
    assert int(mok["crime_criminal"]) == 30
    assert int(mok["population"]) == 1000
    assert abs(float(mok["crime_rate_per_1000"]) - 50.0) < 1e-9
    assert abs(float(mok["crime_criminal_rate_per_1000"]) - 30.0) < 1e-9
    assert abs(float(mok["crime_detectability_pct"]) - 54.3) < 1e-9


def test_attach_crime_features_joins_on_district_casefold() -> None:
    crime = build_crime_by_district([
        {
            "district": "Mokotów",
            "year": 2024,
            "population": 1000,
            "crime_total": 50,
            "crime_criminal": 30,
            "crime_detectability_pct": 54.3,
        },
        {
            "district": "Wola",
            "year": 2024,
            "population": 800,
            "crime_total": 80,
            "crime_criminal": 25,
            "crime_detectability_pct": 46.5,
        },
    ])
    listings = pd.DataFrame({
        "listing_id": ["a", "b", "c"],
        "district": ["mokotów", "WOLA", "unknown"],
        "price_pln": [1.0, 2.0, 3.0],
    })
    out = attach_crime_features(listings, crime)
    for col in CRIME_FEATURE_COLUMNS:
        assert col in out.columns
    assert abs(float(out.loc[0, "crime_rate_per_1000"]) - 50.0) < 1e-9
    assert abs(float(out.loc[1, "crime_rate_per_1000"]) - 100.0) < 1e-9
    assert pd.isna(out.loc[2, "crime_rate_per_1000"])


def test_load_crime_by_district_fixture_json(tmp_path: Path) -> None:
    fixture = tmp_path / "crime_by_district.json"
    fixture.write_text(
        '[{"district":"Bemowo","year":2024,"population":100,'
        '"crime_total":10,"crime_criminal":6,"crime_detectability_pct":49.0}]\n',
        encoding="utf-8",
    )
    out = tmp_path / "crime_by_district.parquet"
    frame = load_crime_by_district(out, fixture_path=fixture)
    assert len(frame) == 1
    assert float(frame.iloc[0]["crime_rate_per_1000"]) == 100.0
    assert out.exists()


def test_model_includes_crime_rate_features() -> None:
    from happy_housing.models.train import CORE_FEATURE_COLUMNS

    assert "crime_rate_per_1000" in CORE_FEATURE_COLUMNS
    assert "crime_criminal_rate_per_1000" in CORE_FEATURE_COLUMNS
    assert "crime_detectability_pct" in CORE_FEATURE_COLUMNS
